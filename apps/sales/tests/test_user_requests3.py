"""Foydalanuvchi so'rovlari (3): kun yakuni majburiy (kechagi kun sanalmaguncha sotuv yo'q), hisobot
eksporti (Excel/CSV), "Sotildi" doim ko'rinadi, "0" bosilganda o'chadi."""

import datetime
import json
from pathlib import Path

import pytest
from django.conf import settings
from django.utils import timezone

from apps.sales.models import DailyClose, Sale
from conftest import INSPECTOR_HOST, SELLER_HOST, set_settings


def _sale_on(shop, seller, days_ago, total=50000):
    s = Sale.objects.create(
        shop=shop, seller=seller, subtotal=total, total=total, payment_type="cash"
    )
    when = timezone.now() - datetime.timedelta(days=days_ago)
    Sale.objects.filter(pk=s.pk).update(created_at=when)
    return s


def _api(client, offline=False):
    body = {
        "items": [{"name": "x", "qty": 1, "unit_price": 20000}],
        "payment_type": "cash",
        "mode": "quick",
    }
    if offline:
        body.update(offline=True, client_uid="u-1", client_ts=timezone.now().isoformat())
    return client.post(
        "/api/sales/", json.dumps(body), content_type="application/json", HTTP_HOST=SELLER_HOST
    )


@pytest.mark.django_db
def test_yesterday_not_closed_blocks_new_sales(sclient, shop, seller):
    _sale_on(shop, seller, 1)
    # Boshqa sahifaga otib yubormaydi: shu joyda "yopiq" ekrani va kun yakuniga tugma
    for url, what in (("/sotuv/", "Sotuv"), ("/skaner/", "Sotuv"), ("/kirim/", "Kirim")):
        r = sclient.get(url, HTTP_HOST=SELLER_HOST)
        html = r.content.decode()
        assert r.status_code == 200 and f"{what} vaqtincha yopiq" in html and "/kun-yakuni/" in html, url
    r = _api(sclient)
    assert r.status_code == 423 and r.json()["redirect"] == "/kun-yakuni/"
    # Oflayn navbatdagi sotuv allaqachon bo'lib o'tgan — rad etilmaydi (yo'qolmasin)
    assert _api(sclient, offline=True).status_code == 201
    home = sclient.get("/", HTTP_HOST=SELLER_HOST).content.decode()
    assert "kun yakuni qilinmagan" in home


@pytest.mark.django_db
def test_closing_yesterday_in_the_morning_unlocks_sales(sclient, shop, seller):
    _sale_on(shop, seller, 1)
    yesterday = timezone.localdate() - datetime.timedelta(days=1)
    page = sclient.get("/kun-yakuni/", HTTP_HOST=SELLER_HOST)
    assert page.context["close_day"] == yesterday and page.context["pending"] == yesterday
    sclient.post("/kun-yakuni/", {"counted_cash": "50 000"}, HTTP_HOST=SELLER_HOST)
    assert DailyClose.objects.filter(shop=shop, date=yesterday).exists()
    assert not DailyClose.objects.filter(shop=shop, date=timezone.localdate()).exists()
    assert "vaqtincha yopiq" not in sclient.get("/sotuv/", HTTP_HOST=SELLER_HOST).content.decode()
    assert _api(sclient).status_code == 201


@pytest.mark.django_db
def test_no_gate_without_recent_sales_or_when_disabled(sclient, shop, seller):
    def locked():
        return "vaqtincha yopiq" in sclient.get("/sotuv/", HTTP_HOST=SELLER_HOST).content.decode()

    assert not locked()  # kecha savdo yo'q
    _sale_on(shop, seller, 10)  # 7 kundan eski — pilotdan oldingi ma'lumot yangi do'konni to'smaydi
    assert not locked()
    _sale_on(shop, seller, 1)
    set_settings(require_daily_close=False)
    assert not locked()


@pytest.mark.django_db
def test_seller_report_export_xlsx_and_csv(sclient, shop, seller):
    _sale_on(shop, seller, 0, total=1500)  # kichik summa ham qatorda
    r = sclient.get("/hisobot/eksport/?fmt=xlsx&davr=14", HTTP_HOST=SELLER_HOST)
    assert r.status_code == 200 and "spreadsheetml" in r["Content-Type"]
    import io

    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb["Kunlar"]
    assert ws.max_row == 15  # sarlavha + 14 kun (bo'sh kunlar ham)
    assert "Umumiy" in wb.sheetnames
    r = sclient.get("/hisobot/eksport/?fmt=csv", HTTP_HOST=SELLER_HOST)
    text = r.content.decode("utf-8")
    assert (
        text.startswith("﻿Sana;") and "1500" in text and r["Content-Disposition"].endswith('.csv"')
    )


@pytest.mark.django_db
def test_inspector_shop_export_scoped(iclient, shop):
    r = iclient.get(f"/dokon/{shop.pk}/eksport/?fmt=csv", HTTP_HOST=INSPECTOR_HOST)
    assert r.status_code == 200 and "Kamera (%)" in r.content.decode("utf-8")
    from apps.geo.models import Market, Region
    from apps.shops.models import Shop

    other = Market.objects.create(
        region=Region.objects.first(), name="Boshqa", latitude=41, longitude=69
    )
    alien = Shop.objects.create(
        market=other, number="9", stir="999999999", latitude=41, longitude=69
    )
    assert iclient.get(f"/dokon/{alien.pk}/eksport/", HTTP_HOST=INSPECTOR_HOST).status_code == 404
    assert (
        iclient.get("/hisobot/eksport/?fmt=csv", HTTP_HOST=INSPECTOR_HOST)
        .content.decode("utf-8")
        .startswith("﻿Bozor;")
    )


def test_sell_button_always_reachable_and_zero_clears():
    base = Path(settings.BASE_DIR)
    for tpl in ("seller/sale.html", "seller/scan.html"):
        src = (base / "templates" / tpl).read_text(encoding="utf-8")
        assert 'class="sale-submit-bar"' in src and "Sotildi · " in src, tpl
    css = (base / "static/css/components.css").read_text(encoding="utf-8")
    assert ".sale-submit-bar{position:sticky" in css
    js = (base / "static/js/components.js").read_text(encoding="utf-8")
    assert "function isZeroVal" in js and "t.__zero" in js


@pytest.mark.django_db
def test_home_opens_even_with_unexplained_zero_sales_day(sclient, shop):
    """Bosh sahifa E'tiroz sahifasiga majburan otib yubormaydi — vazifa kartada turadi."""
    from apps.analytics.models import Alert

    Alert.objects.create(shop=shop, date=timezone.localdate() - datetime.timedelta(days=2),
                         kind=Alert.Kind.ZERO_SALES, level="red", reason="x")
    r = sclient.get("/", HTTP_HOST=SELLER_HOST)
    html = r.content.decode()
    assert r.status_code == 200 and "Savdosiz kun sababini yozing" in html and "/e-tiroz/" in html
