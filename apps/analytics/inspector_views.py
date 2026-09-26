"""Tekshiruvchi interfeysi: dashboard, bozor xaritasi, do'kon sahifasi, signallar, tekshiruv."""

from datetime import timedelta

from django.contrib.auth.decorators import login_required
from django.db.models import Avg
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST

from apps.cameras.models import Camera
from apps.core.dates import day_start, days_between
from apps.geo.models import Market

from .models import Alert, Appeal, DailyScore, Inspection


def _visible_shops(request):
    return request.user.visible_shops().select_related("market", "row", "category")


def _readonly_block(request):
    """Kuzatuvchi (prokuror) yaratish formalarini ochmaydi."""
    if getattr(request.user, "is_prosecutor", False):
        return render(request, "403.html", {
            "heading": "Kuzatuvchi rejimi",
            "message": "Sizning hisobingiz faqat ko'rish uchun: yangi yozuv yaratib bo'lmaydi.",
        }, status=403)
    return None


def _gate_rows(shop):
    from apps.analytics.gate import shop_summary

    return shop_summary(shop)


def _pk(value) -> int:
    """So'rovdan kelgan ID — raqam bo'lmasa 0 ("abc" bilan 500 xato bo'lmasin)."""
    from apps.core.format import to_int

    return to_int(value, 0) or 0




def _level(truth, cfg):
    if truth >= cfg.green_threshold:
        return "green"
    if truth >= cfg.yellow_threshold:
        return "yellow"
    return "red"


def _score_date(shops, today):
    """Bugun hali hisoblanmagan bo'lsa (ertalab / fon vazifasi kutilmoqda) — oxirgi o'lchangan kun.

    Aks holda dashboard va xarita har ertalab bo'm-bo'sh ko'rinadi.
    """
    if DailyScore.objects.filter(shop__in=shops, date=today, measured=True).exists():
        return today
    return (
        DailyScore.objects.filter(shop__in=shops, date__lt=today, measured=True)
        .order_by("-date").values_list("date", flat=True).first()
    ) or today


@login_required
def dashboard(request):
    from apps.core.models import SystemSettings

    cfg = SystemSettings.get_solo()
    shops = _visible_shops(request)
    today = timezone.localdate()

    score_date = _score_date(shops, today)
    latest = {
        s.shop_id: s
        for s in DailyScore.objects.filter(shop__in=shops, date=score_date).select_related(
            "shop", "shop__category", "shop__market"
        )
    }
    # Faqat O'LCHANGAN ballar: ma'lumoti yo'q do'kon 0% bo'lib o'rtachani tushirmasin
    measured = [s for s in latest.values() if s.measured]
    truths = [s.truth_pct for s in measured]
    avg_truth = round(sum(truths) / len(truths)) if truths else None

    alerts = Alert.objects.filter(shop__in=shops).select_related("shop", "shop__market")
    new_alerts = alerts.filter(status=Alert.Status.NEW)

    insp = Inspection.objects.filter(shop__in=shops)
    confirmed = insp.filter(result=Inspection.Result.CONFIRMED).count()
    false_sig = insp.filter(result=Inspection.Result.FALSE).count()
    accuracy = round(confirmed / (confirmed + false_sig) * 100) if (confirmed + false_sig) else None

    # Eng xavfli do'konlar (score_date balli bo'yicha)
    risky = sorted(measured, key=lambda s: s.truth_pct)[:8]

    # Yashirilgan savdo (oxirgi 30 kun) + potensial qo'shimcha soliq
    from django.db.models import Sum

    since = today - timedelta(days=29)
    hidden = (
        DailyScore.objects.filter(shop__in=shops, date__gte=since).aggregate(
            h=Sum("hidden_sales")
        )["h"]
        or 0
    )
    potential_tax = int(hidden * cfg.tax_rate_percent / 100)

    ctx = {
        "shop_count": shops.count(),
        "red_count": new_alerts.filter(level="red").count(),
        "yellow_count": new_alerts.filter(level="yellow").count(),
        "avg_truth": avg_truth,
        "accuracy": accuracy,
        "alerts": new_alerts.order_by("-created_at")[:12],
        "risky": [(s, _level(s.truth_pct, cfg)) for s in risky],
        "today": today,
        "score_date": score_date,
        "hidden_sales": int(hidden),
        "potential_tax": potential_tax,
        "tax_rate": cfg.tax_rate_percent,
    }
    return render(request, "inspector/dashboard.html", ctx)


@login_required
def market_map(request, pk=None):
    from apps.core.models import SystemSettings

    cfg = SystemSettings.get_solo()
    markets = Market.objects.filter(id__in=_visible_shops(request).values("market_id")).distinct()
    market = get_object_or_404(markets, pk=pk) if pk else markets.first()
    if market is None:
        return render(request, "inspector/market_map.html", {"markets": markets, "market": None})

    today = timezone.localdate()
    score_date = _score_date(market.shops.all(), today)
    # Ma'lumotsiz (solishtiruvsiz) do'kon 0% qizil emas — "none" (kulrang) bo'lsin
    scores = {
        s.shop_id: (s.truth_pct if s.has_data else None)
        for s in DailyScore.objects.filter(shop__market=market, date=score_date)
    }

    # Bozor sxemasi: do'konlar QATOR bo'yicha guruhlanadi, har rasta rangli % belgi.
    from collections import defaultdict

    by_row = defaultdict(list)
    counts = {"green": 0, "yellow": 0, "red": 0, "none": 0}
    for shop in market.shops.filter(is_active=True).select_related("row").order_by("number"):
        t = scores.get(shop.id)
        lvl = _level(t, cfg) if t is not None else "none"
        counts[lvl] += 1
        row_label = shop.row.label if shop.row_id else "Boshqa"
        by_row[row_label].append(
            {
                "id": shop.id,
                "number": shop.number,
                "owner": shop.owner_name,
                "truth": t,
                "level": lvl,
            }
        )
    rows_data = [{"label": lbl, "stalls": by_row[lbl]} for lbl in sorted(by_row)]

    return render(
        request,
        "inspector/market_map.html",
        {
            "markets": markets,
            "market": market,
            "rows_data": rows_data,
            "counts": counts,
            "score_date": score_date,
            "today": today,
        },
    )


def _investigation(shop, peers, start, today):
    """Nazorat quroli: chegirma vs bozor, tannarxga yaqin sotuvlar."""
    from django.db.models import Sum

    from apps.sales.models import Sale, SaleItem

    rng = (start, today)

    def discount_pct(shops):
        agg = Sale.objects.filter(shop__in=shops, **days_between("created_at", *rng)).aggregate(
            d=Sum("discount"), s=Sum("subtotal")
        )
        sub = agg["s"] or 0
        return round((agg["d"] or 0) / sub * 100, 1) if sub else 0.0

    # Tannarxga yaqin sotuvlar (narx tannarxdan ≤10% yuqori) — dumping/yashirish belgisi.
    # Ustama savdo turiga bog'liq (elektronikada tabiatan kichik, mevada katta), shuning uchun
    # o'xshash do'konlar (bir bozor + bir tur) ulushi bilan birga beriladi: signal — farq.
    from django.db.models import Count, F, Q

    def near_cost(shops):
        agg = SaleItem.objects.filter(
            sale__shop__in=shops, **days_between("sale__created_at", *rng),
            product__isnull=False, product__buy_price__gt=0,
        ).aggregate(n=Count("id"), near=Count("id", filter=Q(unit_price__lte=F("product__buy_price") * 1.1)))
        return agg["near"] or 0, agg["n"] or 0

    near, total = near_cost([shop])
    pct = round(near / total * 100) if total else 0
    p_near, p_total = near_cost(list(peers)) if peers else (0, 0)
    peer_pct = round(p_near / p_total * 100) if p_total else None
    shop_disc, market_disc = discount_pct([shop]), discount_pct(list(peers) + [shop])
    return {
        "shop_discount": shop_disc,
        "market_discount": market_disc,
        "discount_flag": shop_disc > market_disc + 5,
        "near_cost": near,
        "near_cost_total": total,
        "near_cost_pct": pct,
        "peer_near_cost_pct": peer_pct,
        "near_cost_flag": pct >= (max(30, peer_pct + 20) if peer_pct is not None else 30),
    }


@login_required
def shop_detail(request, pk):
    from apps.core.models import SystemSettings

    cfg = SystemSettings.get_solo()
    shop = get_object_or_404(_visible_shops(request), pk=pk)
    today = timezone.localdate()
    start = today - timedelta(days=13)

    scores = list(DailyScore.objects.filter(shop=shop, date__range=(start, today)).order_by("date"))
    latest = scores[-1] if scores else None
    level = _level(latest.truth_pct, cfg) if latest and latest.measured else "none"

    chart = {
        "labels": [s.date.strftime("%d.%m") for s in scores],
        "entered": [s.entered_sales for s in scores],
        "cash": [s.cash_amount for s in scores],
        # O'lchanmagan kun grafikda uzilish (null), 0% emas
        "truth": [s.truth_pct if s.measured else None for s in scores],
    }

    # O'xshash do'konlar (bir bozor + bir toifa) bilan solishtirish
    peers = shop.similar_shops()
    peer_scores = DailyScore.objects.filter(shop__in=peers, date=today, measured=True)
    _pa = peer_scores.aggregate(a=Avg("truth_pct"))["a"]
    peer_avg = round(_pa) if _pa is not None else None

    part_labels = {
        "cash": "Kassa / deklaratsiya",
        "camera": "Kamera",
        "stock": "Qoldiq",
        "price": "Narx",
    }
    parts = []
    if latest:
        for k, lbl in part_labels.items():
            v = latest.parts.get(k)
            parts.append(
                {
                    "key": k,
                    "label": lbl,
                    "value": v,
                    "level": _level(v, cfg) if v is not None else "none",
                }
            )

    ctx = {
        "shop": shop,
        "latest": latest,
        "level": level,
        "parts": parts,
        "chart": chart,  # json_script o'zi serializatsiya qiladi
        "peer_avg": peer_avg,
        "peer_count": peers.count(),
        "invest": _investigation(shop, peers, start, today),
        "alerts": shop.alerts.order_by("-created_at")[:8],
        "inspections": shop.inspections.select_related("inspector")[:6],
        "appeals": shop.appeals.all()[:5],
        "corrections": shop.corrections.select_related("user")[:8],
        "register_closes": shop.register_closes.order_by("-date")[:10],
        "writeoffs": shop.writeoffs.select_related("seller", "product").order_by("-created_at")[:10],
        "stock_ins": shop.stock_ins.select_related("product").order_by("-created_at")[:10],
        "gate_rows": _gate_rows(shop),
    }
    return render(request, "inspector/shop_detail.html", ctx)


@login_required
def shop_evidence(request, pk):
    """Bitta do'kon uchun chop etsa bo'ladigan dalil to'plami (tekshiruv/akt uchun)."""

    from apps.core.models import SystemSettings

    cfg = SystemSettings.get_solo()
    shop = get_object_or_404(_visible_shops(request), pk=pk)
    today = timezone.localdate()
    start = today - timedelta(days=29)

    scores = list(DailyScore.objects.filter(shop=shop, date__range=(start, today)).order_by("date"))
    latest = scores[-1] if scores else None
    hidden = sum(s.hidden_sales for s in scores)
    potential_tax = int(hidden * cfg.tax_rate_percent / 100)
    peers = shop.similar_shops()

    ctx = {
        "shop": shop,
        "latest": latest,
        "level": _level(latest.truth_pct, cfg) if latest and latest.measured else "none",
        "start": start,
        "today": today,
        "hidden_sales": hidden,
        "potential_tax": potential_tax,
        "tax_rate": cfg.tax_rate_percent,
        "entered_total": sum(s.entered_sales for s in scores),
        "cash_total": sum(s.cash_amount for s in scores),
        "invest": _investigation(shop, peers, start, today),
        "alerts": shop.alerts.order_by("-created_at")[:15],
        "register_closes": shop.register_closes.order_by("-date")[:10],
        "inspections": shop.inspections.select_related("inspector")[:6],
        "now": timezone.now(),
        "inspector": request.user,
    }
    return render(request, "inspector/shop_evidence.html", ctx)


@login_required
def alerts_list(request):
    shops = _visible_shops(request)
    alerts = Alert.objects.filter(shop__in=shops).select_related("shop", "shop__market")
    level = request.GET.get("level")
    status = request.GET.get("status", "")
    if level:
        alerts = alerts.filter(level=level)
    if status:
        alerts = alerts.filter(status=status)
    from apps.core.pagination import paginate

    page = paginate(request, alerts.order_by("-created_at"), per_page=50)
    qs = f"level={level or ''}&status={status}"
    return render(
        request,
        "inspector/alerts.html",
        {
            "alerts": page.object_list,
            "page": page,
            "querystring": qs,
            "level": level or "",
            "status": status,
        },
    )


@login_required
@require_POST
def alert_action(request, pk):
    shops = _visible_shops(request)
    alert = get_object_or_404(Alert, pk=pk, shop__in=shops)
    action = request.POST.get("action")
    labels = {"assign_me": "o'ziga oldi", "dismiss": "e'tiborsiz qoldirdi", "reopen": "qayta ochdi"}
    if action == "assign_me":
        alert.assigned_to = request.user
        alert.status = Alert.Status.ASSIGNED
    elif action == "dismiss":
        alert.status = Alert.Status.DISMISSED
    elif action == "reopen":
        alert.status = Alert.Status.NEW
    alert.save()
    # Audit: kim qaysi signalni nima qildi (jimgina yopish izsiz qolmasin)
    request.audit_detail = (
        f"Signal #{alert.pk} ({alert.get_level_display()}, №{alert.shop.number}, "
        f"{alert.date:%d.%m.%Y}) — {labels.get(action, action)}"
    )
    if request.htmx:
        return render(request, "inspector/_alert_row.html", {"a": alert})
    return redirect("inspector:alerts")


@login_required
def inspection_create(request):
    if (blocked := _readonly_block(request)) is not None:
        return blocked
    shops = _visible_shops(request)
    alert_id = request.GET.get("alert")
    alert = Alert.objects.filter(pk=_pk(alert_id), shop__in=shops).first() if alert_id else None
    preselect = alert.shop if alert else None

    if request.method == "POST":
        from django.contrib import messages as _msg

        from apps.core.format import to_int

        shop = get_object_or_404(shops, pk=_pk(request.POST.get("shop")))
        # Signal FAQAT shu do'konniki bo'lsin (boshqa bozor signalini yopib bo'lmasin)
        post_alert = None
        if request.POST.get("alert"):
            post_alert = Alert.objects.filter(pk=_pk(request.POST.get("alert")), shop=shop).first()
        result = request.POST.get("result")
        if result not in Inspection.Result.values:
            # Natija aniq tanlanishi shart (jim standart ayblov bo'lmasin)
            _msg.error(request, "Tekshiruv natijasini tanlang.")
            return redirect(request.get_full_path())
        fine = to_int(request.POST.get("fine_amount"))
        if fine is not None and fine < 0:
            _msg.error(request, "Jarima manfiy bo'lmasin.")
            return redirect(request.get_full_path())
        photo = request.FILES.get("photo")
        if photo is not None:
            from django.core.exceptions import ValidationError

            from apps.core.media import validate_image_upload

            try:
                validate_image_upload(photo)
            except ValidationError as e:
                _msg.error(request, e.messages[0])
                return redirect(request.get_full_path())
        insp = Inspection.objects.create(
            shop=shop,
            alert=post_alert,
            inspector=request.user,
            result=result,
            act_number=request.POST.get("act_number", "").strip()[:60],
            fine_amount=fine or None,
            notes=request.POST.get("notes", ""),
            photo=photo,
        )
        if insp.alert_id:
            insp.alert.status = Alert.Status.RESOLVED
            insp.alert.save(update_fields=["status"])
        request.audit_detail = (
            f"Tekshiruv #{insp.pk}: №{shop.number} — {insp.get_result_display()}"
            + (f", jarima {fine}" if fine else "")
        )
        from django.contrib import messages

        messages.success(request, "Tekshiruv natijasi saqlandi.")
        # Tasdiqlangan bo'lsa — to'g'ridan-to'g'ri jarima aktiga o'tamiz
        if insp.result == Inspection.Result.CONFIRMED:
            return redirect("inspector:inspection_act", pk=insp.pk)
        return redirect("inspector:inspections")

    return render(
        request,
        "inspector/inspection_form.html",
        {
            "shops": shops,
            "alert": alert,
            "preselect": preselect,
            "results": Inspection.Result.choices,
        },
    )


@login_required
def inspection_act(request, pk):
    """Jarima akti (dalolatnoma) — chop etiladigan rasmiy hujjat, avtomatik to'ldiriladi."""
    from apps.core.models import SystemSettings

    cfg = SystemSettings.get_solo()
    shops = _visible_shops(request)
    insp = get_object_or_404(
        Inspection.objects.select_related("shop", "shop__market", "inspector"),
        pk=pk,
        shop__in=shops,
    )
    shop = insp.shop
    # Rasmiy hujjat O'ZGARMAS bo'lsin: davr va sana tekshiruv kuniga qotiriladi
    # (ilgari "bugun"dan hisoblanib, aktni har chop etganda raqam/sana o'zgarardi).
    act_dt = timezone.localtime(insp.created_at)
    today = act_dt.date()
    start = today - timedelta(days=29)
    scores = list(DailyScore.objects.filter(shop=shop, date__range=(start, today)))
    hidden = sum(s.hidden_sales for s in scores)
    evaded_tax = int(hidden * cfg.tax_rate_percent / 100)
    # Takroriylik: shu aktdan OLDIN tasdiqlangan tekshiruv bo'lganmi
    # (keyinroq yaratilgani eski aktni "takroriy" qilib qo'ymasin)
    repeat = (
        Inspection.objects.filter(
            shop=shop, result=Inspection.Result.CONFIRMED, created_at__lt=insp.created_at
        )
        .exclude(pk=insp.pk)
        .exists()
    )
    from apps.analytics.fines import classify_hidden

    tier = classify_hidden(hidden, repeat=repeat)
    suggested_fine = int(evaded_tax * (1 + cfg.fine_penalty_percent / 100))
    if repeat:
        suggested_fine *= 2  # takroriy — 2 baravar

    changed = []
    if not insp.act_number:  # dalolatnoma raqami — avtomatik
        insp.act_number = f"BN-{today.year}-{insp.pk:05d}"
        changed.append("act_number")
    if insp.fine_amount in (None, 0):  # jarima bo'sh bo'lsa — taxminiy to'ldiramiz
        insp.fine_amount = suggested_fine
        changed.append("fine_amount")
    if changed and not request.user.is_prosecutor:  # kuzatuvchi ko'rishi bazaga yozmaydi
        insp.save(update_fields=changed)

    return render(
        request,
        "inspector/inspection_act.html",
        {
            "insp": insp,
            "shop": shop,
            "start": start,
            "today": today,
            "hidden_sales": hidden,
            "evaded_tax": evaded_tax,
            "suggested_fine": suggested_fine,
            "tax_rate": cfg.tax_rate_percent,
            "penalty_pct": cfg.fine_penalty_percent,
            "tier": tier,
            "repeat": repeat,
            "now": act_dt,  # akt sanasi = tekshiruv vaqti (har ochilganda o'zgarmaydi)
            "inspector": insp.inspector or request.user,
        },
    )


@login_required
def shop_search(request):
    """Do'kon qidirish: raqam yoki STIR bo'yicha. Bitta topilsa — to'g'ridan sahifaga."""
    from django.db.models import Q

    from apps.catalog.models import Product

    q = request.GET.get("q", "").strip()
    results = []
    if q:
        results = list(
            _visible_shops(request).filter(Q(number__icontains=q) | Q(stir__icontains=q))[:50]
        )
        if len(results) == 1:
            return redirect("inspector:shop_detail", pk=results[0].pk)
    # Mahsulot bo'yicha: "M o'lchamli ko'ylak qaysi do'konda bor" — joriy qoldiq bilan
    pq = request.GET.get("p", "").strip()
    psize = request.GET.get("size", "").strip()
    items = []
    if pq or psize:
        qs = Product.objects.filter(shop__in=_visible_shops(request), is_active=True, stock__gt=0)
        if pq:
            qs = qs.filter(Q(name__icontains=pq) | Q(base_name__icontains=pq) | Q(barcode=pq))
        if psize:
            from apps.catalog import variants

            if variants.measure(psize):  # "250g" = "250 g" = "0,25 kg"
                ids = [p.pk for p in qs.exclude(size="").only("pk", "size")
                       if variants.same(p.size, psize)]
                qs = qs.filter(pk__in=ids)
            else:
                qs = qs.filter(size__iexact=variants.normalize(psize))
        items = list(
            qs.select_related("shop", "shop__market").order_by("shop__market__name", "shop__number",
                                                              "name")[:300]
        )
    # Jami — faqat bir xil birlikda (5 kg + 3 dona = "8" ma'nosiz)
    units = {p.unit for p in items}
    return render(
        request,
        "inspector/shop_search.html",
        {"q": q, "results": results, "pq": pq, "psize": psize, "items": items,
         "total_qty": sum(p.stock for p in items) if len(units) == 1 else None,
         "total_unit": items[0].get_unit_display() if len(units) == 1 else ""},
    )


@login_required
def appeals_list(request):
    """Barcha e'tirozlar bir joyda (ilgari faqat do'kon sahifasida — topish qiyin edi)."""
    from apps.core.pagination import paginate

    shops = _visible_shops(request)
    base = Appeal.objects.filter(shop__in=shops)
    status = request.GET.get("status", "new")
    if status not in Appeal.Status.values:
        status = ""
    qs = base.select_related("shop", "shop__market", "alert", "author")
    if status:
        qs = qs.filter(status=status)
    page = paginate(request, qs.order_by("-created_at"), per_page=30)
    return render(
        request,
        "inspector/appeals.html",
        {
            "appeals": page.object_list,
            "page": page,
            "querystring": f"status={status}",
            "status": status,
            "new_count": base.filter(status=Appeal.Status.NEW).count(),
        },
    )


@login_required
def inspections_list(request):
    """Tekshiruvlar va aktlar arxivi — o'tgan dalolatnomani qayta topib chop etish uchun."""
    from django.db.models import Count, Q, Sum

    from apps.core.pagination import paginate

    shops = _visible_shops(request)
    base = Inspection.objects.filter(shop__in=shops)
    result = request.GET.get("result", "")
    if result not in Inspection.Result.values:
        result = ""
    mine = request.GET.get("mine") == "1"
    q = request.GET.get("q", "").strip()
    qs = base.select_related("shop", "shop__market", "inspector", "alert")
    if result:
        qs = qs.filter(result=result)
    if mine:
        qs = qs.filter(inspector=request.user)
    if q:
        qs = qs.filter(
            Q(shop__number__icontains=q) | Q(shop__stir__icontains=q) | Q(act_number__icontains=q)
        )
    totals = base.aggregate(
        total=Count("id"),
        confirmed=Count("id", filter=Q(result=Inspection.Result.CONFIRMED)),
        false=Count("id", filter=Q(result=Inspection.Result.FALSE)),
        fines=Sum("fine_amount", filter=Q(result=Inspection.Result.CONFIRMED)),
    )
    page = paginate(request, qs.order_by("-created_at"), per_page=30)
    return render(
        request,
        "inspector/inspections.html",
        {
            "inspections": page.object_list,
            "page": page,
            "querystring": f"result={result}&mine={'1' if mine else ''}&q={q}",
            "result": result,
            "mine": mine,
            "q": q,
            "totals": totals,
        },
    )


@login_required
@require_POST
def appeal_respond(request, pk):
    """Inspektor e'tirozga javob beradi (qabul/rad + matn)."""
    from django.contrib import messages
    from django.utils.http import url_has_allowed_host_and_scheme

    shops = _visible_shops(request)
    appeal = get_object_or_404(Appeal, pk=pk, shop__in=shops)
    nxt = request.POST.get("next", "")
    if not url_has_allowed_host_and_scheme(nxt, allowed_hosts={request.get_host()}):
        nxt = ""
    back = redirect(nxt) if nxt else redirect("inspector:shop_detail", pk=appeal.shop_id)
    action = request.POST.get("action")
    response = request.POST.get("response", "").strip()[:2000]
    if appeal.status != Appeal.Status.NEW:
        messages.info(request, "Bu e'tirozga allaqachon javob berilgan.")
        return back
    if action == "rejected" and not response:
        # Sababsiz rad etish — sotuvchi nima uchunligini bilmaydi, adolatsiz
        messages.error(request, "Rad etish sababini yozing.")
        return back
    if action in ("accepted", "rejected"):
        request.audit_detail = f"E'tiroz #{appeal.pk} (№{appeal.shop.number}) — {action}"
        appeal.status = action
        appeal.response = response or ("Qabul qilindi." if action == "accepted" else "")
        appeal.save(update_fields=["status", "response"])
        # E'tiroz qabul qilindi = signal asossiz — ochiq signal yopiladi
        if action == "accepted" and appeal.alert_id and appeal.alert.status in (
            Alert.Status.NEW, Alert.Status.ASSIGNED
        ):
            appeal.alert.status = Alert.Status.DISMISSED
            appeal.alert.save(update_fields=["status"])
        from apps.core.models import Notification, notify

        target = appeal.author or getattr(appeal.shop, "staff", None)
        if hasattr(target, "first"):  # staff manager
            target = target.first()
        verdict = "qabul qilindi" if action == "accepted" else "rad etildi"
        notify(
            target,
            Notification.Kind.APPEAL_REPLY,
            f"E'tirozingizga javob: {verdict}",
            body=appeal.response[:200],
            url="/e-tiroz/",
            key=f"appeal:{appeal.id}:{action}",
        )
        messages.success(request, "E'tirozga javob berildi.")
    return back


def _parse_at(request):
    """?sana=YYYY-MM-DD&vaqt=HH:MM → aware datetime (Toshkent). Bo'sh → None (joriy holat)."""
    from datetime import datetime

    raw_d = (request.GET.get("sana") or "").strip()
    if not raw_d:
        return None
    try:
        d = datetime.strptime(raw_d, "%Y-%m-%d").date()
    except ValueError:
        return None
    raw_t = (request.GET.get("vaqt") or "").strip() or "23:59"
    try:
        t = datetime.strptime(raw_t, "%H:%M").time()
    except ValueError:
        t = datetime.strptime("23:59", "%H:%M").time()
    # Daqiqa OXIRIgacha: "14:00" tanlansa 14:00:35 dagi harakat ham kirsin
    t = t.replace(second=59, microsecond=999999)
    at = timezone.make_aware(datetime.combine(d, t))
    return min(at, timezone.now())


@login_required
def inventory(request, pk=None):
    """Ombor — har do'konda qancha mahsulot bor: HOZIR yoki tanlangan SANA/VAQTda.

    Prokuratura ssenariysi: "bozor yondi, o'sha kuni soat 14:00 da qaysi do'konda nima
    qancha bor edi" — javob o'zgarmas tovar harakati jurnalidan (StockMove) olinadi.
    """
    from django.db.models import Count, DecimalField, F, Q, Sum

    from apps.catalog.models import Product
    from apps.sales.services.stock import last_count_at, stock_at

    markets = Market.objects.filter(id__in=_visible_shops(request).values("market_id")).distinct()
    market = get_object_or_404(markets, pk=pk) if pk else markets.first()
    at = _parse_at(request)
    rows = []
    total_value = total_items = 0
    if market:
        shops = list(market.shops.filter(is_active=True).order_by("number"))
        if at is None:
            # BITTA agregat so'rov: har do'kon bo'yicha qiymat/mahsulot soni (N+1 yo'q)
            agg = {
                r["shop"]: (int(r["value"] or 0), r["items"] or 0)
                for r in Product.objects.filter(shop__market=market, is_active=True)
                .values("shop")
                .annotate(
                    value=Sum(F("stock") * F("sell_price"), output_field=DecimalField()),
                    items=Count("id", filter=Q(stock__gt=0)),
                )
            }
        else:
            snap = stock_at(shops, at)
            shop_of = dict(Product.objects.filter(pk__in=snap).values_list("pk", "shop_id"))
            agg = {}
            for pid, (bal, price, _t) in snap.items():
                if bal > 0:
                    v, n = agg.get(shop_of[pid], (0, 0))
                    agg[shop_of[pid]] = (v + int(bal * price), n + 1)
        counts = last_count_at(shops, at or timezone.now())
        for shop in shops:
            value, items = agg.get(shop.id, (0, 0))
            total_value += value
            total_items += items
            rows.append({"shop": shop, "value": value, "items": items,
                         "last_count": counts.get(shop.id)})
        rows.sort(key=lambda r: r["value"], reverse=True)
    return render(
        request,
        "inspector/inventory.html",
        {
            "markets": markets,
            "market": market,
            "rows": rows,
            "total_value": total_value,
            "total_items": total_items,
            "at": at,
            "at_query": (f"?sana={at:%Y-%m-%d}&vaqt={at:%H:%M}" if at else ""),
        },
    )


@login_required
def shop_inventory(request, pk):
    """Bitta do'kon ombori (hozir yoki tanlangan vaqtda) — chop etiladigan ma'lumotnoma.

    Jurnal butunligi (xesh zanjiri) tekshiriladi: yozuv o'zgartirilgan/o'chirilgan bo'lsa ko'rinadi.
    """
    from apps.catalog.models import Product
    from apps.sales.models import StockMove
    from apps.sales.services.stock import last_count_at, stock_at, verify_chain

    shop = get_object_or_404(_visible_shops(request), pk=pk)
    at = _parse_at(request)
    prods = list(Product.objects.filter(shop=shop).order_by("name"))
    if at is None:
        lines = [
            {"p": p, "qty": p.stock, "price": p.sell_price, "last": None}
            for p in prods if p.is_active
        ]
    else:
        snap = stock_at([shop], at)
        by_id = {p.pk: p for p in prods}
        lines = [
            {"p": by_id[pid], "qty": bal, "price": price, "last": t}
            for pid, (bal, price, t) in snap.items()
            if pid in by_id
        ]
    for ln in lines:
        ln["value"] = int(max(ln["qty"], 0) * ln["price"])
    from apps.catalog.sizes import product_order

    lines.sort(key=lambda ln: product_order(ln["p"]))  # model → razmer (36, 37 ... yonma-yon)
    ok, bad_id, n_moves = verify_chain(shop)
    moment = at or timezone.now()
    last_move = (
        StockMove.objects.filter(shop=shop, created_at__lte=moment)
        .order_by("-created_at", "-id").first()
    )
    first_move = StockMove.objects.filter(shop=shop).order_by("created_at").first()
    # Tanlangan kundagi harakatlar (o'sha paytgacha) — nima kirdi/chiqdi
    day_moves = []
    if at is not None:
        day_moves = list(
            StockMove.objects.filter(shop=shop, created_at__gte=day_start(timezone.localtime(at).date()),
                                     created_at__lte=at)
            .select_related("product").order_by("-created_at")[:50]
        )
    from apps.sales.models import DailyClose

    close_photo = (
        DailyClose.objects.filter(shop=shop, date__lte=timezone.localtime(moment).date())
        .exclude(photo="").order_by("-date").first()
    )
    return render(
        request,
        "inspector/shop_inventory.html",
        {
            "shop": shop,
            "close_photo": close_photo,
            "lines": [ln for ln in lines if ln["qty"] != 0],
            "total": sum(ln["value"] for ln in lines),
            "at": at,
            "moment": moment,
            "last_count": last_count_at([shop], moment).get(shop.id),
            "chain_ok": ok,
            "chain_bad": bad_id,
            "chain_n": n_moves,
            "fingerprint": last_move.hash[:16] if last_move else "",
            "journal_start": first_move.created_at if first_move else None,
            "before_journal": bool(at and first_move and at < first_move.created_at),
            "day_moves": day_moves,
            "generated": timezone.now(),
            "inspector": request.user,
        },
    )


@login_required
def cameras_status(request):
    shops = _visible_shops(request)
    cams = Camera.objects.filter(market__in=shops.values("market_id")).select_related(
        "market", "shop"
    )
    return render(
        request,
        "inspector/cameras.html",
        {"cameras": cams, "online": sum(1 for c in cams if c.is_online), "total": cams.count()},
    )


def _report_range(request):
    from datetime import datetime

    today = timezone.localdate()
    try:
        start = datetime.strptime(request.GET["start"], "%Y-%m-%d").date()
    except (KeyError, ValueError):
        start = today - timedelta(days=30)
    try:
        end = datetime.strptime(request.GET["end"], "%Y-%m-%d").date()
    except (KeyError, ValueError):
        end = today
    if start > end:  # teskari tanlansa — almashtiramiz
        start, end = end, start
    if (end - start).days > 366:  # juda katta davr serverni qiynamasin
        start = end - timedelta(days=366)
    return start, end


@login_required
def reports(request):
    shops = _visible_shops(request)
    start, end = _report_range(request)
    scores = DailyScore.objects.filter(shop__in=shops, date__range=(start, end))
    insp = Inspection.objects.filter(shop__in=shops, **days_between("created_at", start, end))
    confirmed = insp.filter(result=Inspection.Result.CONFIRMED).count()
    false_sig = insp.filter(result=Inspection.Result.FALSE).count()
    ctx = {
        "start": start,
        "end": end,
        "entered": sum(s.entered_sales for s in scores),
        "cash": sum(s.cash_amount for s in scores),
        "avg_truth": round(
            scores.filter(measured=True).aggregate(a=Avg("truth_pct"))["a"] or 0
        ),
        "alert_count": Alert.objects.filter(shop__in=shops, date__range=(start, end)).count(),
        "inspection_count": insp.count(),
        "confirmed": confirmed,
        "false_signal": false_sig,
        "accuracy": (
            round(confirmed / (confirmed + false_sig) * 100) if (confirmed + false_sig) else None
        ),
        "fines": sum(i.fine_amount or 0 for i in insp),
    }
    return render(request, "inspector/reports.html", ctx)


@login_required
def statistics(request):
    """Sotuv dinamikasi (grafik) + sotuvchilar statistikasi (jadval, trend bilan)."""
    from collections import defaultdict

    from django.db.models import Count, Q, Sum

    from apps.core.models import SystemSettings

    cfg = SystemSettings.get_solo()
    shops = _visible_shops(request)
    start, end = _report_range(request)

    scores = list(
        DailyScore.objects.filter(shop__in=shops, date__range=(start, end)).select_related("shop")
    )

    # 1) Sotuv dinamikasi — kunlik agregat (barcha ko'rinadigan do'konlar bo'yicha)
    by_date = defaultdict(lambda: {"entered": 0, "cash": 0, "truth": [], "count": 0})
    for s in scores:
        d = by_date[s.date]
        d["entered"] += s.entered_sales
        d["cash"] += s.cash_amount
        if s.measured:
            d["truth"].append(s.truth_pct)
        d["count"] += 1
    days = sorted(by_date)
    dynamics = {
        "labels": [d.strftime("%d.%m") for d in days],
        "entered": [by_date[d]["entered"] for d in days],
        "cash": [by_date[d]["cash"] for d in days],
        "truth": [
            round(sum(by_date[d]["truth"]) / len(by_date[d]["truth"])) if by_date[d]["truth"]
            else None
            for d in days
        ],
    }

    # 2) Sotuvchilar statistikasi — har do'kon: savdo, o'rtacha rostlik, signal, trend
    per_shop = defaultdict(lambda: {"entered": 0, "cash": 0, "truth": []})
    mid = start + (end - start) / 2  # trend: davr ikkiga bo'linadi
    half = defaultdict(lambda: {"a": [], "b": []})  # a=birinchi yarim, b=ikkinchi yarim
    for s in scores:
        p = per_shop[s.shop_id]
        p["entered"] += s.entered_sales
        p["cash"] += s.cash_amount
        if s.measured:
            p["truth"].append(s.truth_pct)
            half[s.shop_id]["b" if s.date >= mid else "a"].append(s.truth_pct)

    alert_counts = dict(
        Alert.objects.filter(shop__in=shops, date__range=(start, end))
        .values_list("shop")
        .annotate(n=Count("id"))
    )
    shop_by_id = {s.id: s for s in shops}

    rows = []
    for sid, p in per_shop.items():
        shop = shop_by_id.get(sid)
        if shop is None:
            continue
        avg_truth = round(sum(p["truth"]) / len(p["truth"])) if p["truth"] else None
        a, b = half[sid]["a"], half[sid]["b"]
        trend = None
        if a and b:
            trend = round(sum(b) / len(b) - sum(a) / len(a))
        rows.append(
            {
                "shop": shop,
                "entered": p["entered"],
                "cash": p["cash"],
                "avg_truth": avg_truth,
                "level": _level(avg_truth, cfg) if avg_truth is not None else "none",
                "alerts": alert_counts.get(sid, 0),
                "trend": trend,
            }
        )
    sort = request.GET.get("sort", "truth")
    keymap = {
        # O'lchanmaganlar oxirida (None 0% deb eng xavfli ko'rinmasin)
        "truth": lambda r: (r["avg_truth"] is None, r["avg_truth"] or 0),
        "entered": lambda r: -r["entered"],
        "alerts": lambda r: -r["alerts"],
    }
    rows.sort(key=keymap.get(sort, keymap["truth"]))

    totals = DailyScore.objects.filter(shop__in=shops, date__range=(start, end)).aggregate(
        e=Sum("entered_sales"),
        c=Sum("cash_amount"),
        t=Avg("truth_pct", filter=Q(measured=True)),
    )
    from apps.core.pagination import paginate

    seller_count = len(rows)
    page = paginate(request, rows, per_page=50)
    ctx = {
        "start": start,
        "end": end,
        "sort": sort,
        "dynamics": dynamics,
        "rows": page.object_list,
        "page": page,
        "querystring": f"start={start:%Y-%m-%d}&end={end:%Y-%m-%d}&sort={sort}",
        "total_entered": int(totals["e"] or 0),
        "total_cash": int(totals["c"] or 0),
        "avg_truth": round(totals["t"] or 0),
        "seller_count": seller_count,
    }
    return render(request, "inspector/statistics.html", ctx)


@login_required
def export_excel(request):
    import openpyxl

    shops = _visible_shops(request)
    start, end = _report_range(request)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hisobot"
    ws.append(
        [
            "Bozor",
            "Do'kon",
            "STIR",
            "Egasi",
            "O'rtacha rostlik %",
            "Kiritilgan (so'm)",
            "Deklaratsiya (so'm)",
        ]
    )
    from django.db.models import Q, Sum

    rows = (
        DailyScore.objects.filter(shop__in=shops, date__range=(start, end))
        .values("shop__market__name", "shop__number", "shop__stir", "shop__owner_name")
        .annotate(
            t=Avg("truth_pct", filter=Q(measured=True)),
            e=Sum("entered_sales"),
            c=Sum("cash_amount"),
        )
    )

    from apps.core.format import excel_row  # har katak formula in'ektsiyasidan himoyalangan

    for r in rows:
        ws.append(
            excel_row([
                r["shop__market__name"],
                r["shop__number"],
                r["shop__stir"],
                r["shop__owner_name"],
                round(r["t"]) if r["t"] is not None else "—",
                int(r["e"] or 0),
                int(r["c"] or 0),
            ])
        )
    from django.http import HttpResponse

    resp = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    resp["Content-Disposition"] = f'attachment; filename="hisobot_{start}_{end}.xlsx"'
    wb.save(resp)
    return resp


# ============================================================================
# FAVQULODDA HOLAT — bozor omborini muhrlash (yong'in, suv toshqini...)
# ============================================================================


def _visible_markets(request):
    return Market.objects.filter(id__in=_visible_shops(request).values("market_id")).distinct()


@login_required
def incidents_list(request):
    from .models import Incident

    items = Incident.objects.filter(market__in=_visible_markets(request)).select_related(
        "market", "created_by"
    )
    return render(request, "inspector/incidents.html", {"incidents": items})


@login_required
def incident_create(request):
    """Hodisani e'lon qilish: bozor + tur + sana/vaqt → shu paytdagi holat MUHRLANADI."""
    if (blocked := _readonly_block(request)) is not None:
        return blocked
    from django.contrib import messages

    from .incidents import declare_incident
    from .models import Incident

    markets = _visible_markets(request)
    if request.method == "POST":
        market = markets.filter(pk=_pk(request.POST.get("market"))).first()
        kind = request.POST.get("kind")
        at = _parse_at_post(request)
        errors = []
        if market is None:
            errors.append("bozor")
        if kind not in Incident.Kind.values:
            errors.append("hodisa turi")
        if at is None:
            errors.append("sana va soat")
        if errors:
            messages.error(request, "Tanlang: " + ", ".join(errors) + ".")
            return redirect("inspector:incident_create")
        inc = declare_incident(market, kind, at, request.POST.get("description", "").strip(),
                               request.user)
        request.audit_detail = (
            f"Favqulodda holat {inc.number}: {inc.get_kind_display()}, {market.name}, "
            f"{timezone.localtime(at):%d.%m.%Y %H:%M}"
        )
        messages.success(request, f"{inc.number} muhrlandi: {inc.shops_count} ta do'kon holati saqlandi.")
        return redirect("inspector:incident_detail", pk=inc.pk)
    return render(request, "inspector/incident_form.html",
                  {"markets": markets, "kinds": Incident.Kind.choices})


def _parse_at_post(request):
    """POST sana (YYYY-MM-DD) + soat (HH:MM) → aware datetime, kelajak bo'lmasin."""
    from datetime import datetime

    try:
        d = datetime.strptime((request.POST.get("sana") or "").strip(), "%Y-%m-%d").date()
        t = datetime.strptime((request.POST.get("vaqt") or "").strip(), "%H:%M").time()
    except ValueError:
        return None
    at = timezone.make_aware(datetime.combine(d, t.replace(second=59, microsecond=999999)))
    return at if at <= timezone.now() + timedelta(minutes=1) else None


@login_required
def incident_detail(request, pk):
    from .incidents import snapshot_intact
    from .models import Incident

    inc = get_object_or_404(
        Incident.objects.select_related("market", "created_by"),
        pk=pk, market__in=_visible_markets(request),
    )
    from django.utils.dateparse import parse_datetime

    shops = [dict(s, last_count_dt=parse_datetime(s["last_count"]) if s.get("last_count") else None)
             for s in inc.snapshot.get("shops", [])]
    q = (request.GET.get("q") or "").strip().lower()
    if q:
        shops = [s for s in shops if q in str(s["number"]).lower() or q in (s["owner"] or "").lower()
                 or q in (s["stir"] or "")]
    return render(
        request,
        "inspector/incident_detail.html",
        {
            "inc": inc,
            "shops": sorted(shops, key=lambda s: -s["value"]),
            "intact": snapshot_intact(inc),
            "print_all": request.GET.get("print") == "1",
            "q": q,
            "generated": timezone.now(),
        },
    )


@login_required
def incident_export(request, pk):
    """Muhrlangan holat Excel'da: Umumiy varaq + har do'kon mahsulotlari bitta jadvalda."""
    import openpyxl
    from django.http import HttpResponse

    from .models import Incident

    inc = get_object_or_404(Incident, pk=pk, market__in=_visible_markets(request))

    from apps.core.format import excel_row  # HAR katak: o'lcham/rang ham sotuvchi erkin matni

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Umumiy"
    ws.append(excel_row([f"{inc.number} — {inc.get_kind_display()}", inc.market.name,
                         timezone.localtime(inc.occurred_at).strftime("%d.%m.%Y %H:%M")]))
    ws.append(["Nazorat xeshi", inc.snapshot_hash])
    ws.append([])
    ws.append(["Do'kon", "Egasi", "STIR", "Telefon", "Mahsulot turlari", "Qiymat (so'm)",
               "Oxirgi sanoq", "Jurnal butun"])
    for s in inc.snapshot.get("shops", []):
        ws.append(excel_row([s["number"], s["owner"], s["stir"], s["phone"],
                             s["items"], s["value"], s["last_count"] or "—",
                             "ha" if s["chain_ok"] else "BUZILGAN"]))
    ws2 = wb.create_sheet("Mahsulotlar")
    ws2.append(["Do'kon", "Mahsulot", "O'lcham", "Rang", "Barkod", "Qoldiq", "Birlik",
                "Narx", "Qiymat"])
    for s in inc.snapshot.get("shops", []):
        for ln in s["lines"]:
            ws2.append(excel_row([s["number"], ln["product"], ln["size"], ln["color"],
                                  ln["barcode"], float(ln["qty"]), ln["unit"], ln["price"],
                                  ln["value"]]))
    resp = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    resp["Content-Disposition"] = f'attachment; filename="{inc.number}.xlsx"'
    wb.save(resp)
    request.audit_action = "export"
    request.audit_detail = f"Favqulodda holat {inc.number} Excel"
    return resp
