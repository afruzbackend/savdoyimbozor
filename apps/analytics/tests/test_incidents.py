"""Favqulodda holat: e'lon, muhr (snapshot + xesh), buzishni aniqlash, eksport."""

from datetime import timedelta
from decimal import Decimal

import pytest
from django.utils import timezone

from apps.analytics.incidents import snapshot_intact
from apps.analytics.models import Incident
from conftest import INSPECTOR_HOST, SELLER_HOST


@pytest.mark.django_db
def test_declare_incident_seals_state(iclient, sclient, shop, product, market):
    product.stock = Decimal("0")
    product.sell_price = 10000
    product.save()
    sclient.post("/kirim/", {"product": product.pk, "quantity": "20", "unit_price": "5000"},
                 HTTP_HOST=SELLER_HOST)
    at = timezone.localtime()
    r = iclient.post("/hodisalar/yangi/", {"market": market.pk, "kind": "fire",
                                            "sana": f"{at:%Y-%m-%d}", "vaqt": f"{at:%H:%M}",
                                            "description": "3-qator"}, HTTP_HOST=INSPECTOR_HOST)
    inc = Incident.objects.get()
    assert r.status_code == 302 and inc.number.startswith("FH-")
    s = next(x for x in inc.snapshot["shops"] if x["id"] == shop.pk)
    assert s["value"] == 200000 and s["lines"][0]["qty"] == "20.000"
    # Hodisadan KEYINGI kirim muhrga ta'sir qilmaydi
    sclient.post("/kirim/", {"product": product.pk, "quantity": "99", "unit_price": "1"},
                 HTTP_HOST=SELLER_HOST)
    inc.refresh_from_db()
    assert next(x for x in inc.snapshot["shops"] if x["id"] == shop.pk)["value"] == 200000
    assert snapshot_intact(inc)
    r = iclient.get(f"/hodisalar/{inc.pk}/", HTTP_HOST=INSPECTOR_HOST)
    assert r.status_code == 200 and r.context["intact"] is True
    r = iclient.get(f"/hodisalar/{inc.pk}/excel/", HTTP_HOST=INSPECTOR_HOST)
    assert r.status_code == 200 and "spreadsheet" in r["Content-Type"]


@pytest.mark.django_db
def test_tampered_snapshot_detected(iclient, market):
    from apps.analytics.incidents import declare_incident

    inc = declare_incident(market, "fire", timezone.now(), "", None)
    snap = dict(inc.snapshot, total_value=999999999)
    Incident.objects.filter(pk=inc.pk).update(snapshot=snap)  # bazada qo'lda o'zgartirish
    inc.refresh_from_db()
    assert snapshot_intact(inc) is False
    with pytest.raises(ValueError):
        inc.snapshot_hash = "x" * 64
        inc.save()


@pytest.mark.django_db
def test_incident_requires_fields_and_no_future(iclient, market):
    iclient.post("/hodisalar/yangi/", {"market": market.pk}, HTTP_HOST=INSPECTOR_HOST)
    future = timezone.localtime() + timedelta(days=2)
    iclient.post("/hodisalar/yangi/", {"market": market.pk, "kind": "fire",
                                       "sana": f"{future:%Y-%m-%d}", "vaqt": "10:00"},
                 HTTP_HOST=INSPECTOR_HOST)
    assert not Incident.objects.exists()


@pytest.mark.django_db
def test_other_market_incident_hidden(iclient, market):
    from apps.analytics.incidents import declare_incident
    from apps.geo.models import Market

    other = Market.objects.create(region=market.region, name="Boshqa bozor")
    inc = declare_incident(other, "fire", timezone.now(), "", None)
    assert iclient.get(f"/hodisalar/{inc.pk}/", HTTP_HOST=INSPECTOR_HOST).status_code == 404


@pytest.mark.django_db
def test_exports_neutralize_excel_formulas(iclient, shop, product, market):
    """Sotuvchi erkin matni (rang, o'lcham, egasi) "=HYPERLINK(...)" bo'lsa — Excel'da formula
    bo'lib ishlamaydi (oldin favqulodda holat eksportida o'lcham/rang himoyasiz edi, "-" esa hech qayerda)."""
    import io

    import openpyxl

    from apps.analytics.incidents import declare_incident
    from apps.analytics.models import DailyScore
    from apps.core.format import excel_safe

    assert excel_safe("-2+3+cmd|' /C calc'!A0").startswith("'") and excel_safe(-5) == -5
    evil = '=HYPERLINK("http://x","bos")'
    from apps.sales.models import StockMove
    from apps.sales.services.stock import record_move

    product.color, product.size = evil, "-1+1"
    product.save()
    record_move(product, StockMove.Kind.IN, delta=Decimal("3"), ref="test")  # muhr jurnaldan oladi
    shop.owner_name = "@SUM(1)"
    shop.save()
    inc = declare_incident(market, "fire", timezone.now(), "", None)
    wb = openpyxl.load_workbook(io.BytesIO(iclient.get(f"/hodisalar/{inc.pk}/excel/",
                                                       HTTP_HOST=INSPECTOR_HOST).content))
    cells = [c.value for row in wb["Mahsulotlar"].iter_rows(min_row=2) for c in row]
    cells += [c.value for row in wb["Umumiy"].iter_rows(min_row=5) for c in row]
    assert "'" + evil in cells and "'-1+1" in cells and "'@SUM(1)" in cells
    assert not any(isinstance(v, str) and v[:1] in "=+-@" for v in cells)
    DailyScore.objects.create(shop=shop, date=timezone.localdate(), truth_pct=50, measured=True)
    wb = openpyxl.load_workbook(io.BytesIO(iclient.get("/hisobot/eksport/", HTTP_HOST=INSPECTOR_HOST).content))
    assert "'@SUM(1)" in [c.value for row in wb.active.iter_rows(min_row=2) for c in row]
